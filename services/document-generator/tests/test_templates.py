from __future__ import annotations

import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, Twips
from docxtpl import DocxTemplate
from lxml import etree as ET

from timdoc_document_generator import (
    MacWordConverter,
    TemplatePreparationError,
    TemplatePreparer,
    UnsupportedConverter,
    WindowsWordConverter,
    default_converter,
    template_environment,
)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


class CopyingConverter:
    def convert(self, source: Path, destination: Path) -> None:
        shutil.copyfile(source, destination)


def _create_form(path: Path, title: str) -> None:
    """A synthetic form with the same run layout as the real Форма 4 / Форма 11 templates."""
    document = Document()
    document.add_heading(title, level=1)
    # Как в настоящем шаблоне: линия «Место составления» набрана жирным стилем абзаца.
    bold_style = document.styles.add_style("Линия", WD_STYLE_TYPE.PARAGRAPH)
    bold_style.font.bold = True
    document.add_paragraph("_" * 27, style=bold_style)
    place = document.add_paragraph()
    place.add_run(" " * 15)
    place.add_run("Место составления").font.size = Pt(8)
    document.add_paragraph(" " * 30 + "«___»____________20     г.")
    document.add_paragraph(" " * 57 + "Дата составления").runs[0].font.size = Pt(8)
    document.add_paragraph(
        "Настоящий Контрольный лист составлен о том, что сервисный центр (Исполнитель)"
    )
    document.add_paragraph("_" * 85)
    document.add_paragraph("(наименование организации)")
    line = document.add_paragraph()
    line.add_run("провел ")
    line.add_run("работы по обслуживанию")
    line.add_run("_" * 57)
    document.add_paragraph("(марка Продукции)")
    line = document.add_paragraph()
    line.add_run("заводской ")
    line.add_run("номер №" + "_" * 76)
    document.add_paragraph("(номер)")
    document.add_paragraph("совместно с Потребителем " + "_" * 86)
    document.add_paragraph("(наименование организации, ИНН)")
    document.add_paragraph("_" * 118).runs[0].font.size = Pt(8)
    document.add_paragraph("(адрес организации)")
    document.add_paragraph("_" * 86)
    document.add_paragraph("(Ф.И.О., должность и телефон контактного лица Потребителя)")
    document.add_paragraph(
        "Представитель сервисного центра (Исполнителя)  Представитель Потребителя"
    )
    positions = document.add_paragraph()
    positions.add_run("_" * 35)
    for _ in range(3):
        positions.add_run().add_tab()
    positions.add_run("_" * 35)
    document.add_paragraph("(Должность)\t\t\t\t\t\t\t(Должность)")
    names = document.add_paragraph()
    names.add_run("     ")
    names.add_run("________")
    names.add_run("     _____________________")
    for _ in range(3):
        names.add_run().add_tab()
    names.add_run("     ________     _____________________")
    document.add_paragraph("М.П.   (подпись) \t\t\t(Ф.И.О.)\t\t\t\tМ.П.   (подпись) \t\t\t(Ф.И.О.)")
    document.save(path)


def _create_ticket(path: Path) -> None:
    document = Document()
    outer = document.add_table(rows=1, cols=2)
    grid = outer._tbl.find(f"{{{W}}}tblGrid")
    for column, width in zip(grid.findall(f"{{{W}}}gridCol"), ("8505", "7248"), strict=True):
        column.set(f"{{{W}}}w", width)
    for cell, width in zip(outer.rows[0].cells, (8505, 7248), strict=True):
        cell.width = Twips(width)
    for outer_cell in outer.rows[0].cells:
        outer_cell.paragraphs[0].add_run("СВЕДЕНИЯ О ВЛАДЕЛЬЦЕ")
        table = outer_cell.add_table(rows=31, cols=16)
        labels = {
            1: "НАИМЕНОВАНИЕ И МАРКА ТЕХНИКИ",
            3: "УНИКАЛЬНЫЙ ИДЕНТИФИКАЦИОННЫЙ НОМЕР (УИН)",
            8: "НАИМЕНОВАНИЕ ОРГАНИЗАЦИИ",
            10: "ИНН",
            12: "ФАМИЛИЯ, ИМЯ, ОТЧЕСТВО",
            14: "СТРАНА / ПОЧТОВЫЙ ИНДЕКС",
            16: "ОБЛАСТЬ / РАЙОН",
            18: "НАСЕЛЕННЫЙ ПУНКТ",
            20: "УЛИЦА / ДОМ",
            22: "АДРЕС ЭЛЕКТРОННОЙ ПОЧТЫ",
            24: "ТЕЛЕФОН/ФАКС / МОБИЛЬНЫЙ ТЕЛЕФОН",
            26: "НАИМЕНОВАНИЕ СЕРВИСНОГО ЦЕНТРА",
            28: "ДАТА ПОСТАНОВКИ НА ГАРАНТИЙНЫЙ УЧЕТ",
        }
        for row, label in labels.items():
            table.cell(row, 1).text = label
        for row in table.rows:
            for cell in row.cells:
                mark = ET.SubElement(cell.paragraphs[0]._p.get_or_add_pPr(), f"{{{W}}}rPr")
                ET.SubElement(mark, f"{{{W}}}rFonts").set(f"{{{W}}}ascii", "Tahoma")
        outer_cell.add_paragraph("СЕРВИСНЫЙ ЦЕНТР")
        signature = outer_cell.add_paragraph()
        for chunk in ("_" * 15, "_" * 20, "_" * 14, "М.П", "."):
            run = signature.add_run(chunk)
            run.bold = True
            run.font.size = Pt(6)
        outer_cell.add_paragraph("ПОДПИСЬ")
        outer_cell.add_paragraph("")
    document.save(path)


def _prepare(tmp_path: Path) -> Path:
    source = tmp_path / "source"
    source.mkdir()
    _create_ticket(source / "Талон Сведения о владельце.docx")
    _create_form(source / "Форма 4 - Контрольный лист.doc", "Форма 4")
    _create_form(source / "Форма 11 - АВР.doc", "Форма 11")
    return TemplatePreparer().prepare(source, tmp_path / "prepared", converter=CopyingConverter())


def _paragraphs(path: Path) -> list[ET._Element]:
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    return root.findall(".//w:p", NS)


def _text(paragraph: ET._Element) -> str:
    parts = []
    for child in paragraph.iter():
        tag = ET.QName(child).localname
        if tag == "t":
            parts.append(child.text or "")
        elif tag == "tab":
            parts.append("\t")
    return "".join(parts)


def _paragraph_containing(path: Path, fragment: str) -> ET._Element:
    return next(paragraph for paragraph in _paragraphs(path) if fragment in _text(paragraph))


def _filled(variable: str, width: int, *, size: int = 22, bold: bool = False) -> str:
    spec = f"{width}, {size}, {bold}"
    return (
        f"{{{{ {variable} | blank_left({spec}) }}}}"
        f"{{{{ {variable} }}}}"
        f"{{{{ {variable} | blank_right({spec}) }}}}"
    )


def _run_properties(paragraph: ET._Element, fragment: str) -> ET._Element:
    for run in paragraph.findall("./w:r", NS):
        if fragment in "".join(text.text or "" for text in run.findall("./w:t", NS)):
            properties = run.find("./w:rPr", NS)
            return ET.Element(f"{{{W}}}rPr") if properties is None else properties
    raise AssertionError(f"run with {fragment!r} not found")


def test_prepares_three_docxtpl_templates_without_modifying_sources(tmp_path: Path) -> None:
    source = tmp_path / "source"
    source.mkdir()
    _create_ticket(source / "Талон Сведения о владельце.docx")
    _create_form(source / "Форма 4 - Контрольный лист.doc", "Форма 4")
    _create_form(source / "Форма 11 - АВР.doc", "Форма 11")
    before = {path.name: path.read_bytes() for path in source.iterdir()}

    destination = TemplatePreparer().prepare(
        source, tmp_path / "prepared", converter=CopyingConverter()
    )

    assert {path.name for path in destination.glob("*.docx")} == {
        "Талон.docx",
        "Форма 4.docx",
        "Форма 11.docx",
    }
    assert {path.name: path.read_bytes() for path in source.iterdir()} == before
    expected_variables = {
        "customer_name",
        "customer_inn",
        "customer_name_inn",
        "customer_address",
        "equipment_name",
        "serial_number",
        "service_center",
        "service_employee",
        "service_employee_position",
        "service_location",
    }
    found: set[str] = set()
    for path in destination.glob("*.docx"):
        found |= DocxTemplate(path).get_undeclared_template_variables(template_environment())
    assert expected_variables <= found

    compatibility = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    for path in destination.glob("*.docx"):
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
        required = set((root.get(f"{{{compatibility}}}Ignorable") or "").split())
        assert required <= {prefix for prefix in root.nsmap if prefix}


def test_form_values_replace_only_the_blank_and_keep_the_sentence(tmp_path: Path) -> None:
    prepared = _prepare(tmp_path)
    form = prepared / "Форма 4.docx"

    assert _text(_paragraph_containing(form, "провел работы")) == (
        "провел работы по обслуживанию" + _filled("equipment_name", 57)
    )
    assert _text(_paragraph_containing(form, "заводской номер")) == (
        "заводской номер №" + _filled("serial_number", 76)
    )
    assert _text(_paragraph_containing(form, "совместно с Потребителем")) == (
        "совместно с Потребителем " + _filled("customer_name_inn", 86)
    )
    # Линия адреса набрана 8pt: добивка остаётся в её кегле, генератор пересчитает ширину.
    assert _text(_paragraph_containing(form, "customer_address")) == _filled(
        "customer_address", 118, size=16
    )
    assert not any("customer_contact" in _text(p) for p in _paragraphs(form))
    contact_label = next(
        i for i, p in enumerate(_paragraphs(form)) if "контактного лица" in _text(p)
    )
    assert _text(_paragraphs(form)[contact_label - 1]) == "_" * 86, "линия контакта не тронута"
    assert _text(_paragraph_containing(form, "service_center")) == _filled("service_center", 85)
    # «Место составления» в шаблоне набрано жирным стилем абзаца: добивка жирная, значение нет.
    assert _text(_paragraph_containing(form, "service_location")) == _filled(
        "service_location", 27, bold=True
    )
    assert "{{" not in _text(_paragraph_containing(form, "Настоящий Контрольный лист"))


def test_form_date_is_left_for_handwriting(tmp_path: Path) -> None:
    prepared = _prepare(tmp_path)
    paragraph = _paragraph_containing(prepared / "Форма 11.docx", "«___»")
    assert _text(paragraph) == " " * 30 + "«___»____________20     г."
    assert not any("document_date" in _text(p) for p in _paragraphs(prepared / "Форма 11.docx"))


def test_form_values_are_never_bold_even_under_a_bold_paragraph_style(tmp_path: Path) -> None:
    prepared = _prepare(tmp_path)
    paragraph = _paragraph_containing(prepared / "Форма 4.docx", "service_location")
    place = _run_properties(paragraph, "{{ service_location }}")
    assert place.find("./w:b", NS).get(f"{{{W}}}val") == "0"
    assert place.find("./w:bCs", NS).get(f"{{{W}}}val") == "0"
    padding = _run_properties(paragraph, "blank_left")
    assert padding.find("./w:b", NS) is None, "добивка наследует жирность стиля линии"


def test_ticket_values_inherit_the_cell_font_from_the_paragraph_mark(tmp_path: Path) -> None:
    prepared = _prepare(tmp_path)
    with zipfile.ZipFile(prepared / "Талон.docx") as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    for paragraph in root.iter(f"{{{W}}}p"):
        mark = paragraph.find("./w:pPr/w:rPr/w:rFonts", NS)
        if mark is None:
            continue
        for run in paragraph.findall("./w:r", NS):
            if "{{" not in "".join(text.text or "" for text in run.findall("./w:t", NS)):
                continue
            fonts = run.find("./w:rPr/w:rFonts", NS)
            assert fonts is not None, _text(paragraph)
            assert fonts.get(f"{{{W}}}ascii") == mark.get(f"{{{W}}}ascii")


def test_form_values_use_body_font_size_and_underline(tmp_path: Path) -> None:
    prepared = _prepare(tmp_path)
    form = prepared / "Форма 4.docx"
    address = _run_properties(
        _paragraph_containing(form, "customer_address"), "{{ customer_address }}"
    )
    assert address.find("./w:sz", NS).get(f"{{{W}}}val") == "22"
    padding = _run_properties(_paragraph_containing(form, "customer_address"), "blank_left")
    assert padding.find("./w:sz", NS).get(f"{{{W}}}val") == "16", "добивка остаётся в кегле линии"
    assert padding.find("./w:u", NS) is None
    assert address.find("./w:szCs", NS).get(f"{{{W}}}val") == "22"
    assert address.find("./w:u", NS).get(f"{{{W}}}val") == "single"
    children = [ET.QName(child).localname for child in address]
    assert children.index("sz") < children.index("u")

    model = _paragraph_containing(form, "equipment_name")
    prefix = _run_properties(model, "провел")
    assert prefix.find("./w:u", NS) is None


def test_form_signature_block_keeps_tabs_and_the_customer_side(tmp_path: Path) -> None:
    prepared = _prepare(tmp_path)
    form = prepared / "Форма 4.docx"

    positions = _text(_paragraph_containing(form, "service_employee_position"))
    assert positions == _filled("service_employee_position", 35) + "\t\t\t" + "_" * 35

    names = _text(_paragraph_containing(form, "service_employee |"))
    assert names == (
        "     ________     " + _filled("service_employee", 21) + "\t\t\t"
        "     ________     _____________________"
    )


def test_ticket_signature_matches_the_filled_reference(tmp_path: Path) -> None:
    prepared = _prepare(tmp_path)
    for paragraph in _paragraphs(prepared / "Талон.docx"):
        if "service_employee" not in _text(paragraph):
            continue
        assert _text(paragraph) == "_____________________ {{ service_employee }} М.П."
        value = _run_properties(paragraph, "service_employee")
        assert value.find("./w:b", NS) is None
        assert value.find("./w:sz", NS) is None
        assert value.find("./w:u", NS).get(f"{{{W}}}val") == "single", "линия идёт под фамилией"
        line = _run_properties(paragraph, "_______________")
        assert line.find("./w:b", NS) is not None


def test_ticket_cells_lose_trailing_empty_paragraphs_to_stay_on_one_page(tmp_path: Path) -> None:
    prepared = _prepare(tmp_path)
    with zipfile.ZipFile(prepared / "Талон.docx") as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    outer_cells = root.find(".//w:body/w:tbl", NS).findall("./w:tr/w:tc", NS)
    assert len(outer_cells) == 2
    for cell in outer_cells:
        last = list(cell)[-1]
        assert ET.QName(last).localname == "p"
        assert _text(last) == "ПОДПИСЬ"


def test_ticket_warranty_date_stays_blank_and_copies_have_equal_width(tmp_path: Path) -> None:
    prepared = _prepare(tmp_path)
    with zipfile.ZipFile(prepared / "Талон.docx") as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    assert b"document_date" not in ET.tostring(root)
    outer = root.find(".//w:body/w:tbl", NS)
    columns = [int(c.get(f"{{{W}}}w")) for c in outer.findall("./w:tblGrid/w:gridCol", NS)]
    assert columns == [7876, 7877]
    cells = [
        c.find("./w:tcPr/w:tcW", NS).get(f"{{{W}}}w") for c in outer.findall("./w:tr/w:tc", NS)
    ]
    assert cells == ["7876", "7877"]


def test_ticket_source_may_be_a_doc_converted_through_word(tmp_path: Path) -> None:
    source = tmp_path / "source"
    source.mkdir()
    _create_ticket(source / "Талон Сведения о владельце.doc")
    _create_form(source / "Форма 4.doc", "Форма 4")
    _create_form(source / "Форма 11.doc", "Форма 11")
    converted: list[str] = []

    class RecordingConverter(CopyingConverter):
        def convert(self, source: Path, destination: Path) -> None:
            converted.append(source.name)
            super().convert(source, destination)

    prepared = TemplatePreparer().prepare(
        source, tmp_path / "prepared", converter=RecordingConverter()
    )
    assert converted == ["Талон Сведения о владельце.doc", "Форма 4.doc", "Форма 11.doc"]
    assert "{{ customer_name }}" in "\n".join(
        _text(p) for p in _paragraphs(prepared / "Талон.docx")
    )


def test_ticket_template_has_fourteen_serial_cells_for_each_printed_copy(tmp_path: Path) -> None:
    prepared = _prepare(tmp_path)
    with zipfile.ZipFile(prepared / "Талон.docx") as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    for index in range(14):
        assert xml.count(f"serial_{index:02d}") == 2


def test_missing_source_form_is_reported(tmp_path: Path) -> None:
    source = tmp_path / "source"
    source.mkdir()
    _create_ticket(source / "Талон.docx")

    with pytest.raises(TemplatePreparationError, match="Форма 4"):
        TemplatePreparer().prepare(source, tmp_path / "prepared", converter=CopyingConverter())


def test_label_without_a_blank_line_above_is_reported(tmp_path: Path) -> None:
    source = tmp_path / "source"
    source.mkdir()
    _create_ticket(source / "Талон.docx")
    _create_form(source / "Форма 11.doc", "Форма 11")
    document = Document()
    document.add_paragraph("Сервисный центр")
    document.add_paragraph("(наименование организации)")
    document.save(source / "Форма 4.doc")

    with pytest.raises(TemplatePreparationError, match="нет линии"):
        TemplatePreparer().prepare(source, tmp_path / "prepared", converter=CopyingConverter())


def test_prepared_templates_are_reused_until_sources_change(tmp_path: Path) -> None:
    source = tmp_path / "source"
    source.mkdir()
    _create_ticket(source / "Талон.docx")
    _create_form(source / "Форма 4.doc", "Форма 4")
    _create_form(source / "Форма 11.doc", "Форма 11")
    prepared = tmp_path / "prepared"

    TemplatePreparer().prepare(source, prepared, converter=CopyingConverter())
    stamp = (prepared / "Форма 4.docx").stat().st_mtime_ns
    TemplatePreparer().prepare(source, prepared, converter=CopyingConverter())
    assert (prepared / "Форма 4.docx").stat().st_mtime_ns == stamp

    _create_form(source / "Форма 4.doc", "Форма 4 (новая)")
    TemplatePreparer().prepare(source, prepared, converter=CopyingConverter())
    assert (prepared / "Форма 4.docx").stat().st_mtime_ns != stamp


def test_default_converter_follows_the_platform() -> None:
    assert isinstance(default_converter("win32"), WindowsWordConverter)
    assert isinstance(default_converter("darwin"), MacWordConverter)
    assert isinstance(default_converter("linux"), UnsupportedConverter)
    with pytest.raises(TemplatePreparationError, match="Windows или macOS"):
        UnsupportedConverter().convert(Path("a.doc"), Path("a.docx"))


@pytest.mark.skipif(sys.platform == "win32", reason="Word COM is the real converter on Windows")
def test_windows_converter_refuses_other_platforms(tmp_path: Path) -> None:
    with pytest.raises(TemplatePreparationError, match="Windows"):
        WindowsWordConverter().convert(tmp_path / "a.doc", tmp_path / "a.docx")


def test_mac_converter_stages_files_inside_the_office_container(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(sys, "platform", "darwin")
    container = tmp_path / "UBF8T346G9.Office"
    container.mkdir()
    calls: list[list[str]] = []

    def runner(command: list[str], **_: Any) -> subprocess.CompletedProcess[str]:
        calls.append(command)
        assert Path(command[3]).read_bytes() == b"doc"
        Path(command[4]).write_bytes(b"PK")
        return subprocess.CompletedProcess(command, 0, "", "")

    source = tmp_path / "Форма 4.doc"
    source.write_bytes(b"doc")
    destination = tmp_path / "out" / "Форма 4.docx"
    MacWordConverter(runner=runner, staging_root=container).convert(source, destination)

    assert calls[0][:2] == ["osascript", "-e"]
    assert 'tell application "Microsoft Word"' in calls[0][2]
    staged_source, staged_result = Path(calls[0][3]), Path(calls[0][4])
    assert staged_source.parent.parent == container / "Timdoc"
    assert staged_source.name == "Форма 4.doc"
    assert staged_result.parent == staged_source.parent
    assert destination.read_bytes() == b"PK"
    assert not staged_source.parent.exists()


def test_mac_converter_reports_word_errors(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(sys, "platform", "darwin")
    container = tmp_path / "UBF8T346G9.Office"
    container.mkdir()

    def failing(command: list[str], **_: Any) -> subprocess.CompletedProcess[str]:
        return subprocess.CompletedProcess(command, 1, "", "execution error: (-1708)")

    def silent(command: list[str], **_: Any) -> subprocess.CompletedProcess[str]:
        return subprocess.CompletedProcess(command, 0, "", "")

    def crashing(command: list[str], **_: Any) -> subprocess.CompletedProcess[str]:
        raise subprocess.TimeoutExpired(command, 1)

    source = tmp_path / "Форма 4.doc"
    source.write_bytes(b"doc")
    with pytest.raises(TemplatePreparationError, match="-1708"):
        MacWordConverter(runner=failing, staging_root=container).convert(
            source, tmp_path / "a.docx"
        )
    with pytest.raises(TemplatePreparationError, match="Word не ответил"):
        MacWordConverter(runner=silent, staging_root=container).convert(source, tmp_path / "a.docx")
    with pytest.raises(TemplatePreparationError, match="Word не смог"):
        MacWordConverter(runner=crashing, staging_root=container).convert(
            source, tmp_path / "a.docx"
        )
    with pytest.raises(TemplatePreparationError, match="Не найден Microsoft Word"):
        MacWordConverter(runner=silent, staging_root=tmp_path / "missing").convert(
            source, tmp_path / "a.docx"
        )
    blocked = tmp_path / "blocked"
    blocked.mkdir()
    (blocked / "Timdoc").write_text("не папка", encoding="utf-8")
    with pytest.raises(TemplatePreparationError, match="Нет доступа к папке Word"):
        MacWordConverter(runner=silent, staging_root=blocked).convert(source, tmp_path / "a.docx")

    monkeypatch.setattr(sys, "platform", "linux")
    with pytest.raises(TemplatePreparationError, match="macOS"):
        MacWordConverter(runner=failing).convert(source, tmp_path / "a.docx")


class FakeWordDocument:
    def __init__(
        self, log: list[tuple[str, object]], *, fail: bool = False, close_fails: bool = False
    ) -> None:
        self.log = log
        self.fail = fail
        self.close_fails = close_fails

    def SaveAs2(self, path: str, **options: object) -> None:
        if self.fail:
            raise RuntimeError("Word is busy")
        self.log.append(("save", path))
        Path(path).write_bytes(b"PK")

    def Close(self, **options: object) -> None:
        self.log.append(("close", options))
        if self.close_fails:
            raise RuntimeError("RPC server unavailable")


class FakeWordApplication:
    def __init__(
        self, log: list[tuple[str, object]], *, fail: bool = False, close_fails: bool = False
    ) -> None:
        self.log = log
        self.fail = fail
        self.close_fails = close_fails
        self.Documents = self

    def Open(self, path: str, **options: object) -> FakeWordDocument:
        self.log.append(("open", path))
        return FakeWordDocument(self.log, fail=self.fail, close_fails=self.close_fails)

    def Quit(self) -> None:
        self.log.append(("quit", None))


def _fake_win32com(monkeypatch: pytest.MonkeyPatch, application: FakeWordApplication) -> None:
    import types

    package = types.ModuleType("win32com")
    client = types.ModuleType("win32com.client")
    client.DispatchEx = lambda name: application  # type: ignore[attr-defined]
    package.client = client  # type: ignore[attr-defined]
    monkeypatch.setitem(sys.modules, "win32com", package)
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    monkeypatch.setattr(sys, "platform", "win32")


def test_windows_converter_drives_word_com_and_always_quits(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    log: list[tuple[str, object]] = []
    _fake_win32com(monkeypatch, FakeWordApplication(log))
    source = tmp_path / "Форма 4.doc"
    source.write_bytes(b"doc")
    destination = tmp_path / "out" / "Форма 4.docx"

    WindowsWordConverter().convert(source, destination)

    assert destination.read_bytes() == b"PK"
    assert [entry[0] for entry in log] == ["open", "save", "close", "quit"]
    assert log[2][1] == {"SaveChanges": False}


def test_windows_converter_reports_word_failures_and_quits(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    log: list[tuple[str, object]] = []
    _fake_win32com(monkeypatch, FakeWordApplication(log, fail=True))
    source = tmp_path / "Форма 4.doc"
    source.write_bytes(b"doc")

    with pytest.raises(TemplatePreparationError, match="Word is busy"):
        WindowsWordConverter().convert(source, tmp_path / "Форма 4.docx")
    assert [entry[0] for entry in log] == ["open", "close", "quit"]


def test_windows_converter_quits_word_even_when_close_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    log: list[tuple[str, object]] = []
    _fake_win32com(monkeypatch, FakeWordApplication(log, close_fails=True))
    source = tmp_path / "Форма 4.doc"
    source.write_bytes(b"doc")
    destination = tmp_path / "Форма 4.docx"

    WindowsWordConverter().convert(source, destination)

    assert destination.read_bytes() == b"PK"
    assert [entry[0] for entry in log] == ["open", "save", "close", "quit"]
