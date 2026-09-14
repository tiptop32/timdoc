from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

import pytest
from docx import Document

from timdoc_contracts import Customer, EquipmentItem, GenerationRequest, ServiceSettings
from timdoc_document_generator import (
    DocumentGenerationError,
    DocumentGenerator,
    blank_left,
    blank_padding,
    blank_right,
    metrics,
)


def _write_template(path: Path, title: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph("Хозяйство: {{ customer_name }} / ИНН {{ customer_inn }}")
    document.add_paragraph("Техника: {{ equipment_name }}")
    document.add_paragraph("Заводской номер: {{ serial_number }}")
    document.add_paragraph("Исполнитель: {{ service_center }} / {{ service_employee }}")
    document.add_paragraph("Дата: {{ document_date }}")
    document.add_paragraph(
        "Адрес: {{ customer_address | blank_left(30) }}{{ customer_address }}"
        "{{ customer_address | blank_right(30) }}"
    )
    document.add_paragraph(
        "Контакт: {{ customer_contact | blank_left(20) }}{{ customer_contact }}"
        "{{ customer_contact | blank_right(20) }}"
    )
    document.add_paragraph(
        "Организация: {{ customer_name_inn | blank_left(10, 22, True) }}{{ customer_name_inn }}"
        "{{ customer_name_inn | blank_right(10, 22, True) }}"
    )
    document.save(path)


@pytest.fixture
def template_directory(tmp_path: Path) -> Path:
    directory = tmp_path / "templates"
    directory.mkdir()
    _write_template(directory / "Талон.docx", "Талон")
    _write_template(directory / "Форма 4.docx", "Форма 4")
    _write_template(directory / "Форма 11.docx", "Форма 11")
    return directory


def _request(template_directory: Path, output_directory: Path) -> GenerationRequest:
    return GenerationRequest(
        customer=Customer(name="ООО «Рощинский»", inn="0268104130"),
        settings=ServiceSettings(employee_name="Абдрахманов Т.М."),
        equipment=[
            EquipmentItem(name="RSM SS-780-13", serial_number="M0S07013002106"),
            EquipmentItem(name="RSM SS-780-13", serial_number="M0S07013002105"),
        ],
        document_date="2026-09-07",
        template_directory=template_directory,
        output_directory=output_directory,
    )


def test_generates_three_documents_per_equipment_and_zip(
    tmp_path: Path, template_directory: Path
) -> None:
    before = {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in template_directory.iterdir()
    }

    result = DocumentGenerator().generate(_request(template_directory, tmp_path / "output"))

    assert len(result.files) == 6
    assert Path(result.archive_path).is_file()
    assert {Path(path).name for path in result.files} == {
        "Талон.docx",
        "Форма 4.docx",
        "Форма 11.docx",
    }
    with zipfile.ZipFile(result.archive_path) as archive:
        assert len(archive.namelist()) == 6
        assert all(name.endswith(".docx") for name in archive.namelist())
    after = {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in template_directory.iterdir()
    }
    assert after == before


def test_generated_document_contains_selected_customer_and_equipment(
    tmp_path: Path, template_directory: Path
) -> None:
    result = DocumentGenerator().generate(_request(template_directory, tmp_path / "output"))
    first = next(
        Path(path)
        for path in result.files
        if "M0S07013002106" in path and path.endswith("Талон.docx")
    )

    text = "\n".join(paragraph.text for paragraph in Document(first).paragraphs)
    assert "ООО «Рощинский»" in text
    assert "0268104130" in text
    assert "RSM SS-780-13" in text
    assert "M0S07013002106" in text
    assert "{{" not in text

    with zipfile.ZipFile(template_directory / "Талон.docx") as source_archive:
        source_parts = {
            name: source_archive.read(name)
            for name in source_archive.namelist()
            if name != "word/document.xml"
        }
    with zipfile.ZipFile(first) as result_archive:
        result_parts = {
            name: result_archive.read(name)
            for name in result_archive.namelist()
            if name != "word/document.xml"
        }
    assert result_parts == source_parts


def test_missing_template_is_a_hard_error(tmp_path: Path, template_directory: Path) -> None:
    (template_directory / "Форма 11.docx").unlink()

    with pytest.raises(DocumentGenerationError, match="Форма 11"):
        DocumentGenerator().generate(_request(template_directory, tmp_path / "output"))


def test_rejects_duplicate_serial_numbers(tmp_path: Path, template_directory: Path) -> None:
    request = _request(template_directory, tmp_path / "output")
    request.equipment.append(request.equipment[0])

    with pytest.raises(DocumentGenerationError, match="повторяется"):
        DocumentGenerator().generate(request)


def test_blank_lines_are_filled_or_left_for_handwriting(
    tmp_path: Path, template_directory: Path
) -> None:
    request = _request(template_directory, tmp_path / "output")
    request.customer.address = "с. Рощинский"
    request.customer.representative_position = "Главный инженер"
    request.customer.representative_name = "Иванов И.И."
    result = DocumentGenerator().generate(request)
    first = next(Path(path) for path in result.files if path.endswith("Форма 4.docx"))

    lines = {
        paragraph.text.split(":")[0]: paragraph.text for paragraph in Document(first).paragraphs
    }
    # Ширина линии сохраняется по метрикам Tahoma: текст плюс подчёркивания = 30 подчёркиваний.
    underscore = metrics.char_width("_")
    assert lines["Адрес"] == "Адрес: _________с. Рощинский_________"
    assert (
        0
        <= 30 * underscore - metrics.text_width(lines["Адрес"].removeprefix("Адрес: "))
        < underscore
    )
    assert lines["Контакт"] == "Контакт: Главный инженер Иванов И.И."
    assert lines["Организация"] == "Организация: ООО «Рощинский», ИНН 0268104130"


def test_blank_padding_keeps_the_physical_line_width() -> None:
    assert blank_padding("", 6) == (6, 0)
    assert blank_padding(None, 3) == (3, 0)
    assert blank_left("", 4) == "____" and blank_right("", 4) == ""
    assert blank_padding("слишком длинное значение", 5) == (0, 0)
    cases = (
        ("Инженер", 35, 22, False),
        ("Абдрахманов Т.М.", 21, 22, False),
        ("RSM SS-780-13", 57, 22, False),
        ("г. Уфа", 27, 22, True),
        ("453100, Республика Башкортостан, с. Рощинский, ул. Ленина, 1", 118, 16, False),
    )
    for value, width, size, bold in cases:
        left, right = blank_padding(value, width, size, bold)
        unit = metrics.char_width("_", bold=bold) * size / 22
        total = (left + right) * unit + metrics.text_width(value)
        assert 0 <= width * unit - total < unit, (value, "не шире и не короче одного знака")
        assert abs(left - right) <= 1
    # Жирная линия шире: значение вытесняет меньше жирных подчёркиваний, чем обычных.
    assert (
        sum(blank_padding("г. Уфа", 27, 22, True)) > sum(blank_padding("г. Уфа", 27, 22, False)) - 2
    )
    assert blank_left("г. Уфа", 27, 22, True) == "_" * 11


def test_tahoma_metrics_cover_russian_documents() -> None:
    assert metrics.char_width("_") == 1118
    assert metrics.char_width("_", bold=True) == 1304
    assert metrics.char_width("Ш") > metrics.char_width("и") > metrics.char_width(".")
    assert metrics.char_width("\u2603") == metrics.char_width("n"), (
        "неизвестный символ = средняя буква"
    )
    assert metrics.underscores_for(0) == 0
    assert metrics.underscores_for(-5) == 0
    assert metrics.underscores_for(metrics.char_width("_") * 3) == 3
    assert metrics.underscores_for(metrics.char_width("_") * 3 - 1) == 2, "только вниз"
