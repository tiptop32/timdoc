from pathlib import Path

import pytest
from docx import Document

from timdoc_contracts import Customer, EquipmentItem, GenerationRequest, ServiceSettings
from timdoc_document_generator import DocumentGenerator


@pytest.mark.eval
def test_every_generated_document_has_no_unresolved_template_markers(tmp_path: Path) -> None:
    templates = tmp_path / "templates"
    templates.mkdir()
    for name in ("Талон.docx", "Форма 4.docx", "Форма 11.docx"):
        document = Document()
        document.add_paragraph("{{ customer_name }} {{ equipment_name }} {{ serial_number }}")
        document.save(templates / name)

    result = DocumentGenerator().generate(
        GenerationRequest(
            customer=Customer(name="ООО «Рощинский»", inn="0268104130"),
            settings=ServiceSettings(),
            equipment=[EquipmentItem(name="RSM SS-780-13", serial_number="M0S07013002106")],
            document_date="2026-09-07",
            template_directory=templates,
            output_directory=tmp_path / "out",
        )
    )

    for path in result.files:
        text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
        assert "{{" not in text
        assert "}}" not in text
